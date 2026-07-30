"""Measured comparisons with openpyxl 3.x on identical workbook data."""

from __future__ import annotations

import html
import io
import math
import os
import platform
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(
    0,
    os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "python"
    ),
)

import openpyxl  # noqa: E402

import mojo_openpyxl  # noqa: E402
from mojo_openpyxl._lib import coordinates, escape  # noqa: E402


def timeit(function, repeat=3):
    best = math.inf
    for _ in range(repeat):
        started = time.perf_counter()
        function()
        best = min(best, time.perf_counter() - started)
    return best


def dense_workbook(module, rows=20_000, columns=10):
    workbook = module.Workbook()
    sheet = workbook.active
    for row in range(1, rows + 1):
        sheet.append([row * column + 0.25 for column in range(1, columns + 1)])
    return workbook


def string_workbook(module, rows=30_000):
    workbook = module.Workbook()
    sheet = workbook.active
    for row in range(rows):
        sheet.append(
            [
                f"row {row}",
                "A&B <tag>",
                " leading and trailing ",
                f"=A{row + 1}&B{row + 1}",
            ]
        )
    return workbook


def save_bytes(workbook):
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def main():
    cases = []

    rows = list(range(1, 200_001))
    columns = [(index % 16384) + 1 for index in range(len(rows))]
    coordinates(rows[:10], columns[:10])
    cases.append(
        (
            "200k cell coordinates",
            lambda: coordinates(rows, columns),
            lambda: [
                f"{openpyxl.utils.get_column_letter(column)}{row}"
                for row, column in zip(rows, columns)
            ],
        )
    )

    source = ('<&> "\' plain UTF-8 café ' * 500_000)
    escape(source[:100])
    cases.append(
        (
            f"XML escape {len(source.encode()) / 1e6:.1f} MB",
            lambda: escape(source),
            lambda: html.escape(source, quote=True),
        )
    )

    ours_dense = dense_workbook(mojo_openpyxl)
    upstream_dense = dense_workbook(openpyxl)
    save_bytes(ours_dense)
    cases.append(
        (
            "save dense 20k x 10",
            lambda: save_bytes(ours_dense),
            lambda: save_bytes(upstream_dense),
        )
    )

    ours_text = string_workbook(mojo_openpyxl)
    upstream_text = string_workbook(openpyxl)
    cases.append(
        (
            "save text/formulas 30k x 4",
            lambda: save_bytes(ours_text),
            lambda: save_bytes(upstream_text),
        )
    )

    with tempfile.NamedTemporaryFile(suffix=".xlsx") as temporary:
        upstream_dense.save(temporary.name)
        payload = Path(temporary.name).read_bytes()
    cases.append(
        (
            "load dense 20k x 10",
            lambda: mojo_openpyxl.load_workbook(io.BytesIO(payload)),
            lambda: openpyxl.load_workbook(io.BytesIO(payload)),
        )
    )

    print(f"Machine: {platform.processor() or platform.machine()}, {platform.system()} {platform.release()}")
    print(f"Python {platform.python_version()}, openpyxl {openpyxl.__version__}")
    print()
    print("| case | mojo-openpyxl | openpyxl/reference | result |")
    print("| --- | ---: | ---: | ---: |")
    for name, ours, reference in cases:
        ours_time = timeit(ours)
        reference_time = timeit(reference)
        ratio = reference_time / ours_time
        word = "faster" if ratio >= 1 else "slower"
        print(
            f"| {name} | {ours_time * 1000:.1f} ms | "
            f"{reference_time * 1000:.1f} ms | {ratio:.2f}x {word} |"
        )


if __name__ == "__main__":
    main()
