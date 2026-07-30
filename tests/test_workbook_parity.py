import datetime as dt
import math
from io import BytesIO

import openpyxl
import pytest

import mojo_openpyxl
from mojo_openpyxl.utils.exceptions import IllegalCharacterError


def values(sheet):
    return list(sheet.iter_rows(values_only=True))


def populate(workbook):
    sheet = workbook.active
    sheet.title = "Data & totals"
    sheet.append(["name", "number", "enabled", "formula", "when"])
    sheet.append(
        [
            'A&B <tag> "quoted"',
            12.5,
            True,
            "=B2*2",
            dt.datetime(2025, 7, 8, 9, 10, 11),
        ]
    )
    sheet["A4"] = " whitespace "
    sheet["B4"] = -7
    second = workbook.create_sheet("Second")
    second["C3"] = "#DIV/0!"
    workbook.active = 1
    return workbook


def test_cell_access_append_and_iteration_match_upstream():
    ours = populate(mojo_openpyxl.Workbook())
    theirs = populate(openpyxl.Workbook())
    assert ours.sheetnames == theirs.sheetnames
    assert ours.active.title == theirs.active.title
    for left, right in zip(ours.worksheets, theirs.worksheets):
        assert values(left) == values(right)
    assert ours["Data & totals"]["A2"].coordinate == theirs["Data & totals"]["A2"].coordinate
    assert ours["Data & totals"]["A1:B2"][1][1].value == 12.5


def test_ours_round_trip_all_supported_types(tmp_path):
    path = tmp_path / "ours.xlsx"
    workbook = populate(mojo_openpyxl.Workbook())
    workbook.save(path)
    loaded = mojo_openpyxl.load_workbook(path)
    assert loaded.sheetnames == workbook.sheetnames
    assert loaded.active.title == "Second"
    assert values(loaded["Data & totals"]) == values(workbook["Data & totals"])
    assert loaded["Second"]["C3"].data_type == "e"


def test_real_openpyxl_reads_our_xlsx(tmp_path):
    path = tmp_path / "ours.xlsx"
    populate(mojo_openpyxl.Workbook()).save(path)
    loaded = openpyxl.load_workbook(path)
    assert loaded.sheetnames == ["Data & totals", "Second"]
    assert loaded["Data & totals"]["A2"].value == 'A&B <tag> "quoted"'
    assert loaded["Data & totals"]["E2"].value == dt.datetime(2025, 7, 8, 9, 10, 11)
    assert loaded["Data & totals"]["D2"].value == "=B2*2"
    assert loaded["Data & totals"]["A4"].value == " whitespace "


def test_ours_reads_real_openpyxl_xlsx(tmp_path):
    path = tmp_path / "upstream.xlsx"
    populate(openpyxl.Workbook()).save(path)
    loaded = mojo_openpyxl.load_workbook(path)
    assert loaded.sheetnames == ["Data & totals", "Second"]
    assert loaded.active.title == "Second"
    assert loaded["Data & totals"]["B2"].value == 12.5
    assert loaded["Data & totals"]["C2"].value is True
    assert loaded["Data & totals"]["D2"].value == "=B2*2"
    assert loaded["Data & totals"]["E2"].value == dt.datetime(2025, 7, 8, 9, 10, 11)


def test_date_datetime_and_time_cross_compatibility(tmp_path):
    values_to_test = [
        dt.date(2024, 1, 2),
        dt.datetime(2024, 1, 2, 3, 4, 5),
        dt.time(3, 4, 5),
    ]
    ours_path = tmp_path / "ours-dates.xlsx"
    ours = mojo_openpyxl.Workbook()
    ours.active.append(values_to_test)
    ours.save(ours_path)
    assert list(openpyxl.load_workbook(ours_path).active.values)[0] == tuple(values_to_test)

    upstream_path = tmp_path / "upstream-dates.xlsx"
    upstream = openpyxl.Workbook()
    upstream.active.append(values_to_test)
    upstream.save(upstream_path)
    loaded = list(mojo_openpyxl.load_workbook(upstream_path).active.values)[0]
    assert loaded == (
        dt.datetime(2024, 1, 2),
        dt.datetime(2024, 1, 2, 3, 4, 5),
        dt.time(3, 4, 5),
    )


def test_data_only_matches_upstream_for_uncached_formula(tmp_path):
    path = tmp_path / "formula.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "=1+2"
    workbook.save(path)
    assert mojo_openpyxl.load_workbook(path, data_only=True).active["A1"].value == (
        openpyxl.load_workbook(path, data_only=True).active["A1"].value
    )


def test_unsupported_streaming_and_rich_modes_are_explicit(tmp_path):
    path = tmp_path / "basic.xlsx"
    openpyxl.Workbook().save(path)
    with pytest.raises(NotImplementedError):
        mojo_openpyxl.load_workbook(path, read_only=True)
    with pytest.raises(NotImplementedError):
        mojo_openpyxl.load_workbook(path, rich_text=True)
    with pytest.raises(NotImplementedError):
        mojo_openpyxl.load_workbook(path, keep_vba=True)


def test_write_only_signature_is_explicitly_buffered():
    workbook = mojo_openpyxl.Workbook(write_only=True)
    assert workbook.active is None
    sheet = workbook.create_sheet()
    sheet.append([1])
    assert sheet["A1"].value == 1


def test_file_like_objects_are_supported():
    stream = BytesIO()
    workbook = mojo_openpyxl.Workbook()
    workbook.active.append([1, "two", False])
    workbook.save(stream)
    stream.seek(0)
    assert values(mojo_openpyxl.load_workbook(stream).active) == [(1, "two", False)]


def test_duplicate_titles_match_upstream():
    ours = mojo_openpyxl.Workbook()
    theirs = openpyxl.Workbook()
    ours.create_sheet("Sheet")
    theirs.create_sheet("Sheet")
    ours.create_sheet("sheet")
    theirs.create_sheet("sheet")
    assert ours.sheetnames == theirs.sheetnames


def test_remove_move_and_lookup():
    workbook = mojo_openpyxl.Workbook()
    second = workbook.create_sheet("Second")
    third = workbook.create_sheet("Third")
    workbook.move_sheet(third, offset=-2)
    assert workbook.sheetnames == ["Third", "Sheet", "Second"]
    workbook.remove(second)
    assert workbook["Third"] is third
    assert workbook.index(third) == 0
    assert list(workbook) == workbook.worksheets
    assert workbook.close() is None
    with pytest.raises(KeyError):
        workbook["missing"]


def test_iteration_views_and_dimensions():
    workbook = mojo_openpyxl.Workbook()
    sheet = workbook.active
    sheet["B2"] = 2
    sheet["C3"] = 3
    assert (sheet.min_row, sheet.max_row, sheet.min_column, sheet.max_column) == (
        2,
        3,
        2,
        3,
    )
    assert [[cell.value for cell in row] for row in sheet.rows] == [
        [None, None, None],
        [None, 2, None],
        [None, None, 3],
    ]
    assert [[cell.value for cell in column] for column in sheet.columns] == [
        [None, None, None],
        [None, 2, None],
        [None, None, 3],
    ]
    assert list(sheet.values) == [
        (None, None, None),
        (None, 2, None),
        (None, None, 3),
    ]
    assert tuple(
        tuple(cell.value for cell in column)
        for column in sheet.iter_cols(min_row=2, max_row=3, min_col=2, max_col=3)
    ) == ((2, None), (None, 3))


def test_illegal_control_character_rejected_like_upstream():
    with pytest.raises(IllegalCharacterError):
        mojo_openpyxl.Workbook().active["A1"] = "bad\x01value"
    with pytest.raises(openpyxl.utils.exceptions.IllegalCharacterError):
        openpyxl.Workbook().active["A1"] = "bad\x01value"


@pytest.mark.parametrize("value", [0, -0.0, 1, -3, 1.25, 1e100, True, False])
def test_numeric_and_boolean_round_trip(value, tmp_path):
    path = tmp_path / "number.xlsx"
    workbook = mojo_openpyxl.Workbook()
    workbook.active["A1"] = value
    workbook.save(path)
    loaded = openpyxl.load_workbook(path).active["A1"].value
    if isinstance(value, float):
        assert loaded == pytest.approx(value)
    else:
        assert loaded == value


def test_non_finite_values_follow_openpyxl_empty_value_behavior(tmp_path):
    path = tmp_path / "nan.xlsx"
    workbook = mojo_openpyxl.Workbook()
    workbook.active.append([math.nan, math.inf, -math.inf])
    workbook.save(path)
    assert values(openpyxl.load_workbook(path).active) == [(None, None, None)]
