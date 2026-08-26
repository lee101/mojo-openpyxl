import itertools
import html

import pytest
from openpyxl.utils import cell as upstream

from mojo_openpyxl import utils
from mojo_openpyxl._lib import addr, coordinates, escape, lib


@pytest.mark.parametrize(
    "column", [1, 2, 25, 26, 27, 52, 53, 702, 703, 16384, 18278]
)
def test_get_column_letter_parity(column):
    assert utils.get_column_letter(column) == upstream.get_column_letter(column)


@pytest.mark.parametrize("column", ["A", "z", "AA", "XFD", "ZZZ"])
def test_column_index_parity(column):
    assert utils.column_index_from_string(column) == upstream.column_index_from_string(column)


@pytest.mark.parametrize("column", [0, -1, 18279])
def test_invalid_column_letter_parity(column):
    with pytest.raises(ValueError):
        utils.get_column_letter(column)
    with pytest.raises(ValueError):
        upstream.get_column_letter(column)


@pytest.mark.parametrize(
    "coordinate", ["A1", "$A$1", "ZZ99", "XFD1048576", "$bc$42"]
)
def test_coordinate_parity(coordinate):
    assert utils.coordinate_from_string(coordinate) == upstream.coordinate_from_string(
        coordinate
    )
    if "$" not in coordinate:
        assert utils.coordinate_to_tuple(coordinate) == upstream.coordinate_to_tuple(coordinate)


@pytest.mark.parametrize("value", ["A1", "$A$1", "A1:B2", "A:B", "1:3", "$A1:B$9"])
def test_range_boundaries_and_absolute_parity(value):
    assert utils.range_boundaries(value) == upstream.range_boundaries(value)
    assert utils.absolute_coordinate(value) == upstream.absolute_coordinate(value)


def test_range_generators_parity():
    assert list(utils.rows_from_range("Y2:AB4")) == list(
        upstream.rows_from_range("Y2:AB4")
    )
    assert list(utils.cols_from_range("Y2:AB4")) == list(
        upstream.cols_from_range("Y2:AB4")
    )


def test_range_to_tuple_parity():
    value = "'Bob''s data'!$A$1:C4"
    assert utils.range_to_tuple(value) == upstream.range_to_tuple(value)


def test_mojo_batch_coordinates_match_upstream():
    rows = [1, 9, 10, 999, 1048576] * 200
    columns = list(itertools.islice(itertools.cycle([1, 26, 27, 702, 703]), len(rows)))
    expected = [
        f"{upstream.get_column_letter(column)}{row}"
        for row, column in zip(rows, columns)
    ]
    assert coordinates(rows, columns) == expected


@pytest.mark.parametrize("length", range(1, 18))
def test_mojo_batch_coordinates_simd_validation_tail(length):
    rows = list(range(1, length + 1))
    columns = list(itertools.islice(itertools.cycle([1, 26, 27]), length))
    expected = [
        f"{upstream.get_column_letter(column)}{row}"
        for row, column in zip(rows, columns)
    ]
    assert coordinates(rows, columns) == expected


@pytest.mark.parametrize(
    ("rows", "columns"),
    [
        ([0], [1]),
        ([1048577], [1]),
        ([1], [0]),
        ([1], [18279]),
        ([[1]], [[1]]),
        ([1.5], [1]),
        ([1], [2.5]),
    ],
)
def test_mojo_batch_coordinates_rejects_unsafe_inputs(rows, columns):
    with pytest.raises((TypeError, ValueError)):
        coordinates(rows, columns)


def test_mojo_xml_escape_is_parseable_and_complete():
    import xml.etree.ElementTree as ET

    text = """ leading & <tag> "quoted" and apostrophe: ' trailing """
    encoded = escape(text)
    assert all(entity in encoded for entity in ("&amp;", "&lt;", "&gt;", "&quot;", "&apos;"))
    assert ET.fromstring(f"<v>{encoded}</v>").text == text


def test_native_escape_rejects_short_destination():
    source = b"&"
    destination = bytearray(4)
    assert lib().mox_escape(addr(source), len(source), addr(destination), len(destination)) == -2


def test_native_sized_escape_rejects_inconsistent_size():
    source = b"&"
    destination = bytearray(6)
    assert (
        lib().mox_escape_sized(
            addr(source), len(source), addr(destination), len(destination), 6
        )
        == -3
    )


@pytest.mark.parametrize("length", range(1, 18))
def test_mojo_xml_escape_simd_tail(length):
    source = ("plain<&>\"'text" * 2)[:length]
    expected = html.escape(source, quote=True).replace("&#x27;", "&apos;")
    assert escape(source) == expected


@pytest.mark.parametrize("size", [32767, 32768])
def test_mojo_coordinates_serial_parallel_threshold(size):
    rows = range(1, size + 1)
    columns = itertools.islice(itertools.cycle([1, 26, 27, 702, 703]), size)
    columns = list(columns)
    result = coordinates(rows, columns)
    expected = [
        f"{upstream.get_column_letter(column)}{row}"
        for row, column in zip(rows, columns)
    ]
    assert result == expected
